from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from sqlalchemy import func

from ..core.config import UPLOAD_DIR, settings
from ..core.security import hash_password
from ..models import Config, Content, Course, Enrollment, RAGFragment, RolePermission, User, uid
from .session import SessionLocal, engine
from ..models import Base

PERMISSIONS = {
    "administrador": ["admin:users", "admin:roles", "admin:config", "admin:audit", "course:read"],
    "docente": ["course:read", "course:write", "content:write", "assessment:write", "analytics:read", "rag:query"],
    "estudiante": ["course:read", "rag:query", "assessment:solve"],
}


def create_schema_for_tests(reset: bool = False) -> None:
    if reset:
        Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)


def seed_demo_data(force: bool = False) -> None:
    with SessionLocal() as db:
        if db.query(User).count() and not force:
            return
        if force:
            return
        admin = User(rol="administrador", nombre="Administrador Monda", correo="admin@monda.local", password_hash=hash_password("Admin123!"), estado="activo")
        teacher = User(rol="docente", nombre="Docente Demo", correo="docente@monda.local", password_hash=hash_password("Docente123!"), estado="activo")
        student = User(rol="estudiante", nombre="Estudiante Demo", correo="estudiante@monda.local", password_hash=hash_password("Estudiante123!"), estado="activo")
        db.add_all([admin, teacher, student])
        db.flush()

        course = Course(docente_id=teacher.id, nombre="Introducción a Sistemas", estado="activo")
        db.add(course)
        db.flush()
        db.add(Enrollment(curso_id=course.id, estudiante_id=student.id, estado="activa"))

        seed_texts = [
            "Un sistema es un conjunto de componentes relacionados que interactúan entre sí para alcanzar un propósito u objetivo común.",
            "Los sistemas reciben entradas, realizan procesos y generan salidas. También pueden incorporar retroalimentación para ajustar su funcionamiento.",
            "En un sistema de información, personas, procesos, datos y tecnología colaboran para capturar, procesar, almacenar y distribuir información útil.",
            "La frontera de un sistema permite distinguirlo de su entorno, mientras que las interfaces permiten el intercambio entre componentes o sistemas.",
        ]
        seed_path = UPLOAD_DIR / "seed_sistemas.docx"
        if not seed_path.exists():
            doc = Document()
            doc.add_heading("Unidad 1 — Sistemas y componentes", 0)
            for text in seed_texts:
                doc.add_paragraph(text)
            doc.save(seed_path)
        content = Content(
            documento_id=uid(),
            curso_id=course.id,
            version=1,
            nombre_original="Unidad 1 — Sistemas y componentes.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            tamano_bytes=seed_path.stat().st_size,
            s3_key="local/seed_sistemas.docx",
            hash_sha256=hashlib.sha256(seed_path.read_bytes()).hexdigest(),
            autor_id=teacher.id,
            estado_indexacion="indexado",
            activo=True,
            modelo_embedding="fallback-lexical-v1",
        )
        db.add(content)
        db.flush()
        for i, text in enumerate(seed_texts, 1):
            db.add(RAGFragment(
                contenido_id=content.id,
                curso_id=course.id,
                ordinal=i,
                texto=text,
                ubicacion=f"párrafo {i}",
                metadata_json={"page": 1, "section": "Unidad 1"},
                embedding=None,
                embedding_model="fallback-lexical-v1",
                embedding_provider="fallback",
                hash_fragmento=hashlib.sha256(text.encode()).hexdigest(),
            ))
        db.add(Config(clave="rag", valor={"topK": settings.rag_top_k, "umbral": settings.rag_similarity_threshold, "mode": settings.rag_mode, "modelo": settings.gemini_llm_model, "ollamaModel": settings.ollama_llm_model}, version=1, actualizado_por=admin.id))
        db.add(Config(clave="alerts", valor={"inactividadDias": 7, "promedioMinimo": 11}, version=1, actualizado_por=admin.id))
        for role, perms in PERMISSIONS.items():
            for permission in perms:
                db.add(RolePermission(rol=role, permiso=permission, version=1))
        db.commit()


def init_test_db(reset: bool = True) -> None:
    create_schema_for_tests(reset=reset)
    seed_demo_data()
